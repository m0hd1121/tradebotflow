#!/usr/bin/env python3
"""Export trade journal to Excel (mirrors EURUSD_Journal_Backtest.xlsx schema)."""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path


def _export(db_path: str, out_path: str, backtest_id: str | None) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl not installed — run: pip install openpyxl")

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    query = "SELECT * FROM trades WHERE 1=1"
    params: list = []
    if backtest_id:
        query += " AND backtest_id = ?"
        params.append(backtest_id)
    else:
        query += " AND is_backtest = 0"
    query += " ORDER BY entry_time"

    rows = conn.execute(query, params).fetchall()
    conn.close()

    if not rows:
        print("No trades found.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trade Journal"

    headers = list(rows[0].keys())
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row[key])

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 30)

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    if rows:
        r_vals = [row["r_used"] or row["modelled_r"] or 0.0 for row in rows]
        wins = [r for r in r_vals if r > 0]
        losses = [r for r in r_vals if r < 0]
        win_rate = len(wins) / len(r_vals) if r_vals else 0
        avg_win = sum(wins) / len(wins) if wins else 0
        avg_loss = abs(sum(losses) / len(losses)) if losses else 1
        expectancy = win_rate * avg_win - (1 - win_rate) * avg_loss

        summary = [
            ("Total Trades", len(r_vals)),
            ("Win Rate", f"{win_rate:.1%}"),
            ("Avg Win (R)", f"{avg_win:.3f}"),
            ("Avg Loss (R)", f"{avg_loss:.3f}"),
            ("Expectancy (R)", f"{expectancy:.3f}"),
            ("Total R", f"{sum(r_vals):.2f}"),
            ("Exported", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ]
        for i, (k, v) in enumerate(summary, 1):
            ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
            ws2.cell(row=i, column=2, value=v)

    wb.save(out_path)
    print(f"Exported {len(rows)} trades to {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export trade journal to Excel")
    parser.add_argument("--db", default="data/tradebotflow.db")
    parser.add_argument("--out", default="export/trade_journal.xlsx")
    parser.add_argument("--backtest-id", default=None,
                        help="Export specific backtest; omit for live trades")
    args = parser.parse_args()

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    _export(args.db, args.out, args.backtest_id)


if __name__ == "__main__":
    main()
